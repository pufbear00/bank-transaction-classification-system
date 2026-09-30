"""Classification Rules CRUD API."""
import io
from typing import List, Optional
from fastapi import APIRouter, Depends, HTTPException, Query, UploadFile, File
from fastapi.responses import StreamingResponse
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select
from openpyxl import Workbook
import pandas as pd

from app.dependencies import get_db
from app.models.classification_rule import ClassificationRule
from app.schemas.classification_rule import (
    ClassificationRuleCreate, ClassificationRuleUpdate, ClassificationRuleResponse,
)

router = APIRouter(prefix="/classification-rules", tags=["classification-rules"])

EXPORT_COLUMNS = [
    ("关键词", "keyword"),
    ("Management", "management"),
    ("Grouping", "grouping"),
    ("Class", "class_name"),
    ("CF Line", "cf_line"),
    ("来源", "source"),
    ("命中次数", "match_count"),
]


@router.get("", response_model=List[ClassificationRuleResponse])
async def list_rules(
    source: Optional[str] = None,
    search: Optional[str] = None,
    db: AsyncSession = Depends(get_db),
):
    q = select(ClassificationRule).order_by(ClassificationRule.match_count.desc())
    if source:
        q = q.where(ClassificationRule.source == source)
    if search:
        q = q.where(
            ClassificationRule.keyword.ilike(f"%{search}%")
            | ClassificationRule.management.ilike(f"%{search}%")
        )
    result = await db.execute(q)
    return [ClassificationRuleResponse.model_validate(r) for r in result.scalars().all()]


@router.get("/export")
async def export_rules(db: AsyncSession = Depends(get_db)):
    """Export all classification rules as XLSX."""
    result = await db.execute(
        select(ClassificationRule).order_by(ClassificationRule.match_count.desc())
    )
    rules = list(result.scalars().all())

    wb = Workbook()
    ws = wb.active
    ws.title = "分类规则"
    for col_idx, (header, _) in enumerate(EXPORT_COLUMNS, 1):
        ws.cell(row=1, column=col_idx, value=header)
    for row_idx, rule in enumerate(rules, 2):
        for col_idx, (_, attr) in enumerate(EXPORT_COLUMNS, 1):
            ws.cell(row=row_idx, column=col_idx, value=getattr(rule, attr, None))

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=classification_rules.xlsx"},
    )


@router.post("/import")
async def import_rules(
    file: UploadFile = File(...),
    db: AsyncSession = Depends(get_db),
):
    """Import classification rules from XLSX. Skips duplicates by keyword."""
    content = await file.read()
    try:
        df = pd.read_excel(io.BytesIO(content))
    except Exception:
        raise HTTPException(status_code=400, detail="无法解析文件，请上传 xlsx 格式")

    # Normalize column names
    col_map = {"关键词": "keyword", "Management": "management", "Grouping": "grouping",
               "Class": "class_name", "CF Line": "cf_line", "来源": "source"}
    df = df.rename(columns=col_map)

    required = ["keyword", "management", "grouping", "class_name"]
    for col in required:
        if col not in df.columns:
            raise HTTPException(status_code=400, detail=f"缺少必要列: {col}")

    # Get existing keywords
    result = await db.execute(select(ClassificationRule.keyword))
    existing_keywords = {r for r in result.scalars().all()}

    added = 0
    skipped = 0
    for _, row in df.iterrows():
        keyword = str(row.get("keyword", "")).strip()
        if not keyword or keyword == "nan":
            continue
        management = str(row.get("management", "")).strip()
        if not management or management == "nan":
            continue

        if keyword in existing_keywords:
            skipped += 1
            continue

        grouping = str(row.get("grouping", "")).strip()
        class_name = str(row.get("class_name", "")).strip()
        cf_line_val = row.get("cf_line")
        cf_line = str(cf_line_val).strip() if cf_line_val and str(cf_line_val) != "nan" else None
        source_val = row.get("source")
        source = str(source_val).strip() if source_val and str(source_val) != "nan" else "import"

        rule = ClassificationRule(
            keyword=keyword,
            management=management,
            grouping=grouping if grouping != "nan" else "",
            class_name=class_name if class_name != "nan" else "",
            cf_line=cf_line,
            source=source,
        )
        db.add(rule)
        existing_keywords.add(keyword)
        added += 1

    await db.flush()
    return {"added": added, "skipped": skipped}


@router.post("", response_model=ClassificationRuleResponse)
async def create_rule(data: ClassificationRuleCreate, db: AsyncSession = Depends(get_db)):
    rule = ClassificationRule(**data.model_dump())
    db.add(rule)
    await db.flush()
    return ClassificationRuleResponse.model_validate(rule)


@router.put("/{rule_id}", response_model=ClassificationRuleResponse)
async def update_rule(
    rule_id: int, data: ClassificationRuleUpdate, db: AsyncSession = Depends(get_db)
):
    result = await db.execute(select(ClassificationRule).where(ClassificationRule.id == rule_id))
    rule = result.scalar_one_or_none()
    if not rule:
        raise HTTPException(status_code=404, detail="Rule not found")

    for k, v in data.model_dump(exclude_unset=True).items():
        setattr(rule, k, v)

    return ClassificationRuleResponse.model_validate(rule)


@router.delete("/{rule_id}")
async def delete_rule(rule_id: int, db: AsyncSession = Depends(get_db)):
    result = await db.execute(select(ClassificationRule).where(ClassificationRule.id == rule_id))
    rule = result.scalar_one_or_none()
    if not rule:
        raise HTTPException(status_code=404, detail="Rule not found")

    await db.delete(rule)
    return {"deleted": True}
