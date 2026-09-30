"""Export service: generate XLSX from transactions."""
import io
from typing import List
from openpyxl import Workbook

from app.models.transaction import Transaction


EXPORT_COLUMNS = [
    ("Entity", "entity"),
    ("Account", "account"),
    ("Counterpart", "counterpart"),
    ("Period", "period"),
    ("Abstract", "abstract"),
    ("CCY", "ccy"),
    ("CF", "cf"),
    ("Bal", "bal"),
    ("E.R", "exchange_rate"),
    ("CF in CNY", "cf_in_cny"),
    ("Note", "note"),
    ("Management", "management"),
    ("Grouping", "grouping"),
    ("Class", "class_name"),
    ("CF Line", "cf_line"),
    ("Auto", "classification_method"),
]


def export_transactions_xlsx(transactions: List[Transaction]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Transactions"

    # Header row
    for col_idx, (header, _) in enumerate(EXPORT_COLUMNS, 1):
        ws.cell(row=1, column=col_idx, value=header)

    # Data rows
    for row_idx, txn in enumerate(transactions, 2):
        for col_idx, (_, attr) in enumerate(EXPORT_COLUMNS, 1):
            val = getattr(txn, attr, None)
            if hasattr(val, "isoformat"):
                val = val.isoformat()
            ws.cell(row=row_idx, column=col_idx, value=val)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()
