"""FS for Investor — report generation service."""

import io
import math
from collections import defaultdict
from datetime import date as date_type
from typing import Optional

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from sqlalchemy import select, func, extract, literal_column
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.transaction import Transaction
from app.services.fs_mapping import (
    FSLine,
    PROFIT_LOSS_LINES,
    CASH_FLOW_LINES,
    MAPPED_CF_LINES,
)


def _safe(v) -> float:
    if v is None:
        return 0.0
    f = float(v)
    return 0.0 if (math.isnan(f) or math.isinf(f)) else f


MonthMap = dict[str, float]  # {"2026-01": 1234.56, ...}


class FSReportService:
    def __init__(self, db: AsyncSession):
        self.db = db

    # ─── public ───────────────────────────────
    async def generate(self, year: int, entity: Optional[str] = None):
        months = [f"{year}-{m:02d}" for m in range(1, 13)]
        month_expr = func.strftime("%Y-%m", Transaction.period)

        base_cond = [extract("year", Transaction.period) == year]
        if entity:
            base_cond.append(Transaction.entity == entity)

        # ── 1. Aggregate by grouping+management+month ──
        q_grp = (
            select(
                Transaction.grouping,
                Transaction.management,
                Transaction.import_batch,
                month_expr.label("month"),
                func.sum(Transaction.cf_in_cny).label("total"),
            )
            .where(*base_cond, Transaction.management.isnot(None))
            .group_by(
                Transaction.grouping,
                Transaction.management,
                Transaction.import_batch,
                month_expr,
            )
        )
        rows = (await self.db.execute(q_grp)).all()

        # Build lookup dicts
        by_grouping: dict[str, MonthMap] = defaultdict(lambda: defaultdict(float))
        by_management: dict[str, MonthMap] = defaultdict(lambda: defaultdict(float))
        by_grouping_ex_mgmt: dict[tuple[str, ...], MonthMap] = {}  # (grouping, *excl) → months
        by_amort_neg: MonthMap = defaultdict(float)
        by_finance_plus_interest: MonthMap = defaultdict(float)

        for r in rows:
            grp = r.grouping or ""
            mgmt = r.management or ""
            batch = r.import_batch or ""
            m = r.month
            v = _safe(r.total)

            by_grouping[grp][m] += v
            by_management[mgmt][m] += v

            # Amortization negative cf
            if batch == "amortization" and v < 0:
                by_amort_neg[m] += v

        # ── 2. Aggregate by cf_line+month for cash flow ──
        q_cf = (
            select(
                Transaction.cf_line,
                month_expr.label("month"),
                func.sum(Transaction.cf_in_cny).label("total"),
            )
            .where(*base_cond, Transaction.cf_line.isnot(None))
            .group_by(Transaction.cf_line, month_expr)
        )
        cf_rows = (await self.db.execute(q_cf)).all()
        by_cf_line: dict[str, MonthMap] = defaultdict(lambda: defaultdict(float))
        for r in cf_rows:
            by_cf_line[r.cf_line or ""][r.month] += _safe(r.total)

        # ── 3. FX residual: total cf (excl 99_内部往来) minus all mapped cf_lines ──
        q_total = (
            select(
                month_expr.label("month"),
                func.sum(Transaction.cf_in_cny).label("total"),
            )
            .where(*base_cond, Transaction.class_name != "99_内部往来")
            .group_by(month_expr)
        )
        total_rows = (await self.db.execute(q_total)).all()
        total_by_month: MonthMap = {r.month: _safe(r.total) for r in total_rows}

        mapped_total_by_month: MonthMap = defaultdict(float)
        for cl, mdata in by_cf_line.items():
            if cl in MAPPED_CF_LINES:
                for m, v in mdata.items():
                    mapped_total_by_month[m] += v

        fx_by_month: MonthMap = {}
        for m in months:
            fx_by_month[m] = total_by_month.get(m, 0) - mapped_total_by_month.get(m, 0)

        # ── 4. Opening balance: sum of all cf_in_cny before Jan 1 of year (excl 内部往来) ──
        q_opening = (
            select(func.sum(Transaction.cf_in_cny).label("total"))
            .where(
                Transaction.period < date_type(year, 1, 1),
                Transaction.class_name != "99_内部往来",
            )
        )
        if entity:
            q_opening = q_opening.where(Transaction.entity == entity)
        opening_result = (await self.db.execute(q_opening)).scalar()
        opening_balance = _safe(opening_result)

        # ── 5. Build line values ──
        def _resolve_line(line: FSLine, computed: dict[str, MonthMap]) -> MonthMap:
            result: MonthMap = {m: 0.0 for m in months}

            if line.fixed_zero:
                return result

            # Special: FX residual
            if line.key == "cf_fx":
                return {m: fx_by_month.get(m, 0) for m in months}

            # Special: opening balance (same for all months = cumulative before year)
            if line.key == "cf_opening_bal":
                return {m: opening_balance for m in months}

            # Data source lines
            if line.source_import_batch:
                # amortization with negative cf
                if line.source_cf_negative:
                    return {m: by_amort_neg.get(m, 0) for m in months}

            if line.source_field and line.source_values:
                if line.source_field == "grouping":
                    for gval in line.source_values:
                        for m in months:
                            result[m] += by_grouping[gval].get(m, 0)
                    # Exclude specific managements
                    if line.exclude_management:
                        for mgmt in line.exclude_management:
                            for m in months:
                                result[m] -= by_management[mgmt].get(m, 0)
                    # Include additional managements (from other groupings)
                    if line.include_management:
                        for mgmt in line.include_management:
                            for m in months:
                                result[m] += by_management[mgmt].get(m, 0)
                elif line.source_field == "management":
                    for mval in line.source_values:
                        for m in months:
                            result[m] += by_management[mval].get(m, 0)
                elif line.source_field == "cf_line":
                    for cval in line.source_values:
                        for m in months:
                            result[m] += by_cf_line[cval].get(m, 0)
                return result

            # Formula lines
            if line.formula_add or line.formula_sub:
                for m in months:
                    for ref in line.formula_add:
                        if ref in computed:
                            result[m] += computed[ref].get(m, 0)
                    for ref in line.formula_sub:
                        if ref in computed:
                            result[m] -= computed[ref].get(m, 0)
                return result

            return result

        def _build_statement(lines: list[FSLine]):
            computed: dict[str, MonthMap] = {}
            # Two-pass: first resolve data-source lines, then formula lines
            # This handles cases where a subtotal appears before its children in display order
            for line in lines:
                if not (line.formula_add or line.formula_sub) or line.key in ("cf_fx", "cf_opening_bal"):
                    computed[line.key] = _resolve_line(line, computed)
            for line in lines:
                if line.key not in computed:
                    computed[line.key] = _resolve_line(line, computed)

            output = []
            for line in lines:
                mdata = computed[line.key]
                total = sum(mdata.values())

                # Opening balance total is not a sum — it's just the value
                if line.key == "cf_opening_bal":
                    total = opening_balance

                drilldown = None
                if line.drilldown_field and line.drilldown_value:
                    drilldown = {line.drilldown_field: line.drilldown_value}

                output.append({
                    "key": line.key,
                    "label": line.label,
                    "indent": line.indent,
                    "style": line.style,
                    "months": {m: round(mdata[m], 2) for m in months},
                    "total": round(total, 2),
                    "drilldown": drilldown,
                })
            return output

        pl_lines = _build_statement(PROFIT_LOSS_LINES)
        cf_lines = _build_statement(CASH_FLOW_LINES)

        return {
            "year": year,
            "months": months,
            "statements": [
                {"name": "利润表", "lines": pl_lines},
                {"name": "现金流量表", "lines": cf_lines},
            ],
        }

    # ─── XLSX export ──────────────────────────
    async def export_xlsx(self, year: int, entity: Optional[str] = None) -> bytes:
        data = await self.generate(year, entity)
        months = data["months"]

        wb = Workbook()

        MONTH_LABELS = {
            "01": "1月", "02": "2月", "03": "3月", "04": "4月",
            "05": "5月", "06": "6月", "07": "7月", "08": "8月",
            "09": "9月", "10": "10月", "11": "11月", "12": "12月",
        }

        # Styles
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        section_font = Font(bold=True)
        section_fill = PatternFill(start_color="D6DCE4", end_color="D6DCE4", fill_type="solid")
        subtotal_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
        total_font = Font(bold=True)
        total_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
        num_fmt = "#,##0.00"

        for stmt in data["statements"]:
            ws = wb.create_sheet(title=stmt["name"])

            # Header row
            ws.cell(row=1, column=1, value=stmt["name"]).font = header_font
            ws.cell(row=1, column=1).fill = header_fill
            for ci, m in enumerate(months, 2):
                cell = ws.cell(row=1, column=ci, value=MONTH_LABELS.get(m[5:], m))
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center")
            total_col = len(months) + 2
            cell = ws.cell(row=1, column=total_col, value="全年合计")
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")

            row_idx = 2
            for line in stmt["lines"]:
                style = line["style"]
                indent_str = "    " * line["indent"]
                ws.cell(row=row_idx, column=1, value=f"{indent_str}{line['label']}")

                for ci, m in enumerate(months, 2):
                    c = ws.cell(row=row_idx, column=ci, value=line["months"].get(m, 0))
                    c.number_format = num_fmt

                c = ws.cell(row=row_idx, column=total_col, value=line["total"])
                c.number_format = num_fmt

                # Apply style
                if style == "header":
                    for ci in range(1, total_col + 1):
                        ws.cell(row=row_idx, column=ci).font = section_font
                        ws.cell(row=row_idx, column=ci).fill = section_fill
                elif style == "subtotal":
                    for ci in range(1, total_col + 1):
                        ws.cell(row=row_idx, column=ci).font = Font(bold=True)
                        ws.cell(row=row_idx, column=ci).fill = subtotal_fill
                elif style == "total":
                    for ci in range(1, total_col + 1):
                        ws.cell(row=row_idx, column=ci).font = total_font
                        ws.cell(row=row_idx, column=ci).fill = total_fill

                row_idx += 1

            # Column widths
            ws.column_dimensions["A"].width = 42
            for ci in range(2, total_col + 1):
                ws.column_dimensions[ws.cell(row=1, column=ci).column_letter].width = 14

        # Remove default sheet
        if "Sheet" in wb.sheetnames:
            del wb["Sheet"]

        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()
